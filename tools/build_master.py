# -*- coding: utf-8 -*-
"""入試マスタ（master.xlsx）を作る。
usage: python3 -I build_master.py <raw.json> <haiten.json> <out.xlsx> [haiten_full.json]
haiten_full.json があれば、受験科目・配点の列も埋める。
"""
import sys, json, re, hashlib, collections
sys.path.insert(0, __import__('os').path.dirname(__file__))
from rules import general, extra_general, kyotest, GROUPS, FAC_FILTER, DEPT_FILTER, D, R, ABBR
import added_private as AP
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

raw, haiten_p, out = sys.argv[1:4]
full_p = sys.argv[4] if len(sys.argv) > 4 else None
border_p = sys.argv[5] if len(sys.argv) > 5 else None
GROUP_OF = {u: g for g, us in GROUPS.items() for u in us}
METHOD_HINT = ('方式', '日程', '全学', '学部', '統一', '英語', '技能', '科目', '教科', '型', 'コア', 'プラス', '期',
               '共通', '併用', '利用', '地域枠', 'ＴＥＡＰ', '多面的')

rows = []
def add(**k):
    rows.append(k)

def clean_rank(s):
    s = (s or '').replace('\n', ' ').strip()
    if '設定' in s:
        return '設定なし'
    m = re.match(r'([\d.]+～[\d.]+)', s)
    return m.group(1) if m else s

# ---------- 1. Kei-Net 入試日一覧 ----------
seen = set()
unmatched = collections.Counter()
for x in json.load(open(raw)):
    u, rest, date = x['uni'], x['rest'], x['date']
    note = []
    free = '[自]' in rest
    mm = re.search(r'<(.+?)>', rest)
    if mm:
        note.append(mm.group(1))
    rest = re.sub(r'<.+?>', '', rest).replace('[自]', '').replace('[指]', '').strip()
    p = rest.split('-')
    if len(p) >= 3:
        f, g, m = p[0], p[1], '-'.join(p[2:])
    elif len(p) == 2:
        if any(h in p[1] for h in METHOD_HINT):
            f, g, m = p[0], '', p[1]
        else:
            f, g, m = p[0], p[1], '一般'
    else:
        f, g, m = p[0], '', '一般'
    kind = '共テ併用' if '(テ)' in m else '一般'
    m = m.replace('(テ)', '')
    key = (u, f, g, m, date)
    if key in seen:
        continue
    seen.add(key)
    if free:
        note.insert(0, '試験日自由選択（他日程でも受験可）')
    r = general(u, f, g, m, date)
    if r is None:
        unmatched[(u, f, m)] += 1
        r = R('', '', '', '', conf='要確認')
    if mm and mm.group(1) == '2次':
        fin = re.search(r'最終発表(\d+/\d+)', r['note'] or '')
        if fin:
            r = dict(r, ann=D(fin.group(1)))
    if r['note']:
        note.append(r['note'])
    add(u=u, f=f, g=g, m=m, kind=kind, date=date, rank=clean_rank(x['rank']), r=r, note='／'.join(note), src2='Kei-Net入試日一覧')

# ---------- 2/3. 配点データ（方式一覧）から ----------
H = json.load(open(haiten_p))
def is_kyo(r):
    c, s = r['共通テ総点'], r['二次総点']
    return c not in (None, '', '0', 0) and s in (None, '', 0, '0')

keinet_unis = {x['uni'] for x in json.load(open(raw))}
seen2 = set()
for h in H:
    u = h['大学名']
    if u not in GROUP_OF or GROUP_OF[u] == '国公立' or h['kind'] != '私' or u in AP.D:
        continue
    f, g = h['学部名'] or '', h['学科名'] or ''
    if u in FAC_FILTER and f not in FAC_FILTER[u]:
        continue
    if u in DEPT_FILTER and not any(k in g for k in DEPT_FILTER[u]):
        continue
    m = h['日程方式名'] if h['日程方式名'] not in (None, 'None') else ''
    sub = h['枝番名称'] if h['枝番名称'] not in (None, 'None') else ''
    if is_kyo(h):
        mname = (m or '共テ利用') + (f'（{sub}）' if sub else '')
        k = ('K', u, f, g, mname)
        if k in seen2:
            continue
        seen2.add(k)
        r = kyotest(u, f, g, m)
        if r is None:
            unmatched[(u, f, 'K:' + m)] += 1
            r = R('', '', '', '', conf='要確認')
        add(u=u, f=f, g=g, m=mname, kind='共テ利用', date='', rank='', r=r, note=r['note'], src2='Kei-Net科目・配点')
    elif u not in keinet_unis:
        mname = (m or '一般') + (f'（{sub}）' if sub else '')
        if '地域枠' in mname:
            continue
        k = ('G', u, f, g, mname)
        if k in seen2:
            continue
        seen2.add(k)
        res = extra_general(u, f, g, m, sub)
        if res is None:
            unmatched[(u, f, 'G:' + m)] += 1
            continue
        kind = '共テ併用' if '共通テスト併用' in ''.join(rr['note'] for _, rr in res) else '一般'
        for d, r in res:
            add(u=u, f=f, g=g, m=mname, kind=kind, date=D(d), rank='', r=r, note=r['note'], src2='Kei-Net科目・配点')

# ---------- 4. 国公立 ----------
seen3 = set()
for h in H:
    u = h['大学名']
    if GROUP_OF.get(u) != '国公立' or h['kind'] != '国':
        continue
    f, g, m = h['学部名'] or '', h['学科名'] or '', h['日程方式名'] or ''
    if m in ('', 'None') and u == '東京':
        m = '前期'
    if m not in ('前期', '中期', '後期'):
        continue
    k = (u, f, g, m)
    if k in seen3:
        continue
    seen3.add(k)
    date = D({'前期': '2/25', '中期': '3/8', '後期': '3/12'}[m])
    r = R('1/25', '2/3', '', '', note={'前期': '前期は2/25から（2日間以上の学部あり）。合格発表は例年3/6〜3/10頃',
                                      '中期': '公立中期は3/8から。合格発表は例年3/20頃',
                                      '後期': '後期は3/12から。合格発表は例年3/20〜3/23頃'}[m] + ('。実技あり・日程が異なる。要項で確認' if u == '東京芸術' else '') + '。出願期間・発表日は募集要項（11月頃公表）で確認',
          conf='目安', src='国公立大学共通日程')
    add(u=u, f=f, g=g, m=m + '日程', kind='国公立', date=date, rank='', r=r, note=r['note'], src2='Kei-Net科目・配点')

# ---------- 5. 追加の私立（マナビジョンの日程） ----------
def facs_of(u):
    seen = []
    for h in H:
        if h['大学名'] == u and h['kind'] == '私':
            f = h['学部名'] or ''
            if f and f not in seen and f not in AP.EXCLUDE.get(u, []) and f not in ('医', '歯', '薬'):
                seen.append(f)
    return seen
for u, items in list(AP.D.items()) + list(AP.EXTRA.items()):
    allf = facs_of(u)
    for it in items:
        targets = allf if it['facs'] == ['全'] else it['facs']
        for f in targets:
            fac, dep = (('現代教養', f) if u == '東京女子' else tuple(f.split('|')) if '|' in f else (f, ''))
            r = R('', '', '', '', conf='確定', src=AP.SRC)
            r.update(open=D(it['open']) if '/' in it['open'] else it['open'], close=it['close'], ann=it['ann'], proc=it['proc'])
            for d in (it['dates'] or ['']):
                add(u=u, f=fac, g=dep, m=it['met'], kind=it['kind'], date=d, rank='', r=dict(r),
                    note=it['note'] + ('。' if it['note'] else '') + '科目・配点は要項で確認', src2=AP.SRC, added=True)

# ---------- ID・並べ替え ----------
order_g = list(GROUPS)
order_u = [u for us in GROUPS.values() for u in us]
rows.sort(key=lambda z: (order_g.index(GROUP_OF[z['u']]), order_u.index(z['u']), z['kind'] == '共テ利用', z['f'], z['g'], z['date'] or '9', z['m']))
ids = set()
for z in rows:
    base = '|'.join([z['u'], z['f'], z['g'], z['m'], z['kind'], z['date']])
    i = hashlib.md5(base.encode()).hexdigest()[:8]
    while i in ids:
        i = hashlib.md5((base + i).encode()).hexdigest()[:8]
    ids.add(i)
    z['id'] = i

# ---------- 受験科目・配点 ----------
nsub = collections.Counter()
if full_p:
    import haiten_match as HM
    from subjects import summarize
    HM.load(json.load(open(full_p)))
    for z in rows:
        d, how = HM.find(z['u'], z['f'], z['g'], z['m'], z['kind'])
        z['sub'] = summarize(d) if (d and not z.get('added')) else {}
        nsub['あり' if d else 'なし'] += 1
print('科目データ', nsub)

# ---------- ボーダー（河合塾） ----------
nb = collections.Counter()
if border_p and full_p:
    BRD = {b['大学コード(10桁)']: b for b in json.load(open(border_p))}
    def rk(b):
        if b.get('２次・私大ボーダーランク') == 'BF': return 'BF'
        try: x = float(b.get('２次・私大ボーダーランク(偏差値下限値)'))
        except (TypeError, ValueError): return ''
        return f'{x:.1f}～{x + 2.4:.1f}'
    byfac = collections.defaultdict(list)
    for b in BRD.values():
        byfac[(b['大学名'], (b['学部名'] or '').split('（')[0])].append(b)
    for z in rows:
        if z.get('added'):
            bs = byfac.get((z['u'], z['f'].split('（')[0]), [])
            bs = [b for b in bs if b['学科名'] == z['g']] or bs
            if z['kind'] == '共テ利用':
                v = sorted(int(b['共通テストボーダーライン(得点率)']) for b in bs if str(b.get('共通テストボーダーライン(得点率)')) not in ('None', ''))
                if v: z['note'] += '。共テボーダー（学部内の学科・方式）' + (f'{v[0]}%' if v[0] == v[-1] else f'{v[0]}〜{v[-1]}%'); nb['追加・共テ'] += 1
            else:
                v = sorted(float(b['２次・私大ボーダーランク(偏差値下限値)']) for b in bs if str(b.get('２次・私大ボーダーランク(偏差値下限値)')) not in ('--', 'None', ''))
                if v: z['rank'] = f'{v[0]:.1f}～{v[-1] + 2.4:.1f}'; nb['追加・偏差値'] += 1
            continue
        d, how = HM.find(z['u'], z['f'], z['g'], z['m'], z['kind'])
        b = BRD.get(d.get('１０桁コード')) if d else None
        if not b: nb['なし'] += 1; continue
        r2 = rk(b)
        if r2: z['rank'] = r2
        p = b.get('共通テストボーダーライン(得点率)')
        if str(p) not in ('None', '') and z['kind'] != '一般': z['bd'] = str(p)
        nb['あり'] += 1
print('ボーダー', nb)

# ---------- 書き出し ----------
wb = Workbook()
ws = wb.active
ws.title = '入試日程'
cols = ['ID', '区分', '大学', '学部', '学科', '方式', '種別', '試験日', '出願開始', '出願締切', '合格発表', '手続締切',
        'ボーダー偏差値', '共テボーダー', '共テ科目', '共テ配点', '個別科目', '個別配点', '科目要件', '備考', '確度', '出典']
ws.append(cols)
for z in rows:
    r = z['r']
    ws.append([z['id'], GROUP_OF[z['u']], z['u'], z['f'], z['g'], z['m'], z['kind'], z['date'], r['open'], r['close'],
               r['ann'], r['proc'], z['rank'], z.get('bd', ''), z.get('sub', {}).get('kt', ''), z.get('sub', {}).get('kp', ''),
               z.get('sub', {}).get('nt', ''), z.get('sub', {}).get('np', ''), z.get('sub', {}).get('req', ''),
               z['note'], r['conf'], r['src'] or z['src2']])
widths = [10, 12, 14, 18, 26, 22, 9, 11, 11, 11, 11, 11, 12, 11, 44, 8, 40, 8, 30, 50, 14, 24]
hdr = PatternFill('solid', fgColor='1F3A5F')
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = hdr
ws.freeze_panes = 'D2'
ws.auto_filter.ref = ws.dimensions

w2 = wb.create_sheet('大学')
w2.append(['大学', '表示名', '略称', '区分', '国私'])
for g, us in GROUPS.items():
    for u in us:
        disp = u if u.endswith('大学') else (u + '大学' if u not in ('慶應義塾',) else '慶應義塾大学')
        if u == '東京慈恵会医科':
            disp = '東京慈恵会医科大学'
        w2.append([u, disp, ABBR.get(u, u), g, '国公立' if g == '国公立' else '私立'])
for i, w in enumerate([16, 22, 10, 16, 8], 1):
    w2.column_dimensions[get_column_letter(i)].width = w
    w2.cell(row=1, column=i).font = Font(bold=True)

w3 = wb.create_sheet('設定')
w3.append(['項目', '値', '説明'])
w3.append(['年度', '2027', '入学年度（2027年4月入学）'])
w3.append(['共通テスト1日目', '2027-01-16', ''])
w3.append(['共通テスト2日目', '2027-01-17', ''])
w3.append(['カレンダー開始', '2026-12-01', '受験カレンダーに表示する最初の月'])
w3.append(['カレンダー終了', '2027-03-31', ''])
w3.append(['データ更新日', '2026-10-16', '画面に表示されます'])
w3.append(['注意文', '日程は2026年10月時点の公表情報です。出願前に必ず各大学の募集要項で確認してください。', '画面上部に表示'])
for i, w in enumerate([18, 60, 30], 1):
    w3.column_dimensions[get_column_letter(i)].width = w
    w3.cell(row=1, column=i).font = Font(bold=True)

from openpyxl.worksheet.datavalidation import DataValidation
# 入力規則（プルダウン）
nrow = ws.max_row + 2000
dv_kind = DataValidation(type='list', formula1='"一般,共テ併用,共テ利用,国公立"', allow_blank=True)
dv_conf = DataValidation(type='list', formula1='"確定,目安,前年参考,要確認"', allow_blank=True, showErrorMessage=False)
dv_grp = DataValidation(type='list', formula1="=大学!$D$2:$D$60", allow_blank=True, showErrorMessage=False)
for dv, col in [(dv_kind, 'G'), (dv_conf, 'U'), (dv_grp, 'B')]:
    ws.add_data_validation(dv); dv.add(f'{col}2:{col}{nrow}')

w4 = wb.create_sheet('使い方')
for line in [
    'このファイルが受験校プランナーの大学データです（GitHub の data/master.xlsx）。このファイルを差し替えるだけで、ページの大学データが更新されます。',
    '',
    '■ 年次更新の流れ（だれが担当しても同じ手順です）',
    '1. 教員ページ「大学データ」タブの「更新用Excelをダウンロード」でこのファイルを取得する。',
    '2. 「AIへの依頼文」シートの文章と一緒に、このファイルを外部のAI（ClaudeやChatGPTなど）に渡して更新してもらう。人が手で直してもよい。',
    '3. 更新後のファイルを、教員ページ「大学データ」タブの「更新したExcelをチェック」に入れる。エラーがあれば直す。「この端末で試しに表示」で見た目も確認できる。',
    '4. GitHub の nu1school/jukenko を開き、data フォルダ →「Add file」→「Upload files」で、ファイル名を master.xlsx にしてアップロード →「Commit changes」。1〜2分で公開ページに反映される。',
    '5. Microsoft Forms の回答を「すべての回答を削除」で空にする（新年度の場合）。',
    '',
    '■ シート',
    '「入試日程」：1行＝1つの試験（試験日ごと）。共テ利用・国公立も同じ表。列の意味は「記入ルール」シート。',
    '「大学」：大学の並び順・表示名・略称・区分（早慶上理など）。大学を増やすときはここにも1行足す。',
    '「設定」：年度・共通テスト日・カレンダー期間・データ更新日・注意文。年度更新のたびに書き換える。',
    '「AIへの依頼文」：外部AIに渡す文章。',
]:
    w4.append([line])
w4.column_dimensions['A'].width = 130

w5 = wb.create_sheet('記入ルール')
w5.append(['列', '必須', '書き方', '例'])
for r in [
    ['ID', '', '既存行は変更しない（生徒の登録と結びついている）。新しい行は空欄でよい（自動で振られる）。年度が変わって全行作り直すときは全部空欄でもよい。', 'a1b2c3d4'],
    ['区分', '○', '大学シートの区分から選ぶ（プルダウン）。', 'GMARCH'],
    ['大学', '○', '「大学」を付けない短い名前。大学シートの「大学」列と同じ表記。', '明治'],
    ['学部', '○', '「学部」を付けない。', '理工'],
    ['学科', '', '「学科」を付けない。学部全体の方式なら空欄でよい。', '数学'],
    ['方式', '○', '大学の呼び方どおり。', '全学部統一'],
    ['種別', '○', '一般／共テ併用／共テ利用／国公立 のいずれか（プルダウン）。', '一般'],
    ['試験日', '共テ利用以外は○', 'yyyy-mm-dd。試験日が複数ある方式は、日ごとに1行ずつ。共テ利用は空欄。', '2027-02-05'],
    ['出願開始', '', 'yyyy-mm-dd', '2027-01-07'],
    ['出願締切', '', 'yyyy-mm-dd（Web出願の締切。郵送必着などは備考へ）', '2027-01-18'],
    ['合格発表', '', 'yyyy-mm-dd', '2027-02-15'],
    ['手続締切', '', 'yyyy-mm-dd（入学金の第1次締切）', '2027-02-24'],
    ['ボーダー偏差値', '', '河合塾のボーダー偏差値（帯）など。不明なら空欄。旧名「偏差値帯」でも読める。', '60.0～62.4'],
    ['共テボーダー', '', '共通テストのボーダー得点率（%）。数字だけ書く。共テ利用・共テ併用・国公立の行に。', '78'],
    ['共テ科目', '', '共通テストで使う科目と配点。〔 〕は「どれか選ぶ」、〈 〉は教科をまたいで選ぶ。「・」で区切る。', '英200・数ⅠA・ⅡBC200・〔物/化/生〕200'],
    ['共テ配点', '', '共通テストの合計点（数字）。', '600'],
    ['個別科目', '', '大学の個別試験（一般方式の試験・国公立の2次）の科目と配点。書き方は共テ科目と同じ。面接・小論文・英語資格もここに。', '英100・数ⅠⅡⅢABC200・〔物/化/生〕100'],
    ['個別配点', '', '個別試験の合計点（数字）。', '400'],
    ['科目要件', '', '「受験科目で絞り込み」用。空欄なら絞り込みでも常に表示。書き方：テ:（共テ） 個:（個別）を「 | 」で区切り、科目を空白で区切る。どれか選ぶ科目は「/」、2つ選ぶなら「*2」。使える記号：英 国 数ⅠA 数ⅡBC 数Ⅲ 物 化 生 地学 基礎 日 世 地理 倫 政 情（* は科目以外で代えられる枠）。', 'テ:英 数ⅠA 数ⅡBC 物/化/生 | 個:英 数Ⅲ 物/化*2'],
    ['備考', '', '試験日が選べる、2次試験がある、など。', '試験日自由選択'],
    ['確度', '', '確定（その年度の公表値）／目安／前年参考／要確認。', '確定'],
    ['出典', '', '調べた資料の名前やURL。', '大学公式サイト 入試日程'],
    ['分野', '', '任意。理工・情報／医療・看護／薬学／国際・語学／経済・経営／法・政治／文・教育・心理。空欄なら自動判定。', '理工・情報'],
]:
    w5.append(r)
for i, w in enumerate([12, 16, 90, 22], 1):
    w5.column_dimensions[get_column_letter(i)].width = w
    w5.cell(row=1, column=i).font = Font(bold=True)

w6 = wb.create_sheet('AIへの依頼文')
prompt = """添付のExcel（受験校プランナーの大学データ）を、【 新しい年度：例 2028年度入試（2028年4月入学）】用に更新してください。

■ 守ること
・シートの構成と列の名前・順番は変えない。「入試日程」「大学」「設定」シートを更新し、それ以外のシートはそのまま残す。
・「入試日程」は1行＝1つの試験。試験日が複数ある方式は日ごとに行を分ける。共テ利用は試験日を空欄にする。
・日付はすべて yyyy-mm-dd の文字列で書く（例：2028-02-05）。
・種別は「一般」「共テ併用」「共テ利用」「国公立」のどれか。
・ID列：方式・学部・試験日が変わらない行は前年のIDをそのまま残してよい。新しく作った行のIDは空欄にする。
・対象大学は「大学」シートにある大学だけ。大学・学部・方式の新設や廃止があれば行を足す・消す。
・出願開始・出願締切・合格発表・手続締切は、各大学の公式の入試日程（募集要項・入試ガイド）や、河合塾Kei-Net、ベネッセ マナビジョン、旺文社パスナビで調べる。
・調べた日付には「確度」に「確定」、出典に資料名かURLを書く。まだ発表されていない日付は推測で埋めず空欄にし、確度を「要確認」にする。前年の日付を参考に入れる場合は確度を「前年参考」にする。
・国公立の前期・後期は共通日程で入れ、確度は「目安」。
・ボーダー偏差値・共テボーダー（得点率%、数字のみ）は、業者や河合塾などの資料があれば入れる。なければ空欄。
・共テ科目・共テ配点・個別科目・個別配点・科目要件は、入試方式が変わった行だけ直す（書き方は「記入ルール」シート）。新しい行で分からなければ空欄でよい。
・「設定」シートの年度・共通テスト1日目・2日目・カレンダー開始・カレンダー終了・データ更新日を新しい年度の値にする。
・最後に、確度が「要確認」「前年参考」の行がどの大学に何行あるかを一覧で報告する。

■ 返してほしいもの
・更新したExcelファイル（ファイル名 master.xlsx）
・変更点の要約（追加・削除した方式、日付を確認できなかった大学）"""
for line in prompt.split('\n'):
    w6.append([line])
w6.column_dimensions['A'].width = 130
wb.save(out)

print('rows', len(rows))
print(collections.Counter(z['kind'] for z in rows))
print(collections.Counter(z['r']['conf'] for z in rows))
print('unmatched:')
for k, n in unmatched.most_common(80):
    print(' ', k, n)
